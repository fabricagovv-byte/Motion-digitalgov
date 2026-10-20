# Gera o timbrado DigitalGov (em branco) e a proposta de restauração para Ibiá (DOCX)
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY=RGBColor(0x0A,0x16,0x33); BLUE=RGBColor(0x0E,0x7C,0xB8); GREY=RGBColor(0x5A,0x67,0x84)
LOGO='../digitalgov/assets/logo_digitalgov.png'

# faixa do cabeçalho: logo à esquerda, mosaico à direita
W,H=2400,300
im=Image.new('RGBA',(W,H),(255,255,255,0)); d=ImageDraw.Draw(im)
lg=Image.open(LOGO).convert('RGBA'); lg=lg.resize((int(lg.width*230/lg.height),230)); im.paste(lg,(0,30),lg)
cols=['#1BB8EA','#F5A623','#8DC63F','#E8336D','#263050']
s=56; g=12
grid=[(0,0,0),(1,0,1),(2,0,0),(0,1,2),(1,1,3),(0,2,4),(2,1,0)]
for cx,cy,ci in grid:
    x=W-(cx+1)*(s+g); y=40+cy*(s+g)
    d.rounded_rectangle([x,y,x+s,y+s],8,fill=cols[ci])
im.save('header.png')

def border(p,side='top',color='1BB8EA',sz=18):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); e=OxmlElement(f'w:{side}')
    for k,v in (('val','single'),('sz',str(sz)),('space','6'),('color',color)): e.set(qn('w:'+k),v)
    b.append(e); pPr.append(b)

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)

def base():
    doc=Document()
    st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(11); st.font.color.rgb=NAVY
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(2.2); sec.top_margin=Cm(3.9); sec.bottom_margin=Cm(2.6)
    sec.header_distance=Cm(1); sec.footer_distance=Cm(1)
    hp=sec.header.paragraphs[0]; hp.add_run().add_picture('header.png',width=Cm(16.6))
    border(hp,'bottom','263050',8)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; border(fp,'top','1BB8EA',18)
    r=fp.add_run('DigitalGov · Tecnologia a serviço da sua empresa\n'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=NAVY
    r=fp.add_run('Warley, Diretor · (31) 99827-7323 · digitalgovadm@gmail.com · Alvorada de Minas – MG'); r.font.size=Pt(8.5); r.font.color.rgb=GREY
    return doc

def para(doc,text='',bold=False,size=11,color=NAVY,align=None,after=6,italic=False):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.12
    if align: p.alignment=align
    if text:
        r=p.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size); r.font.color.rgb=color
    return p

def h(doc,text):
    p=para(doc,text.upper(),bold=True,size=11.5,color=BLUE,after=4); p.paragraph_format.space_before=Pt(7)

# 1) timbrado em branco
doc=base()
para(doc,'Alvorada de Minas – MG, ____ de ______________ de 20___.',align=WD_ALIGN_PARAGRAPH.RIGHT,after=24)
para(doc,'À')
para(doc,'[Destinatário]',bold=True,after=2)
para(doc,'[Órgão / Empresa]',after=20)
para(doc,'Assunto: [assunto]',bold=True,after=16)
para(doc,'[Escreva aqui o texto do documento.]')
doc.save('../../deliverables/timbrado_ibia/DigitalGov_Timbrado_em_branco.docx')

# 2) proposta para Ibiá
doc=base()
para(doc,'Alvorada de Minas – MG, ____ de outubro de 2026.',align=WD_ALIGN_PARAGRAPH.RIGHT,after=10)
para(doc,'À')
para(doc,'Prefeitura Municipal de Ibiá – MG',bold=True,after=2)
para(doc,'A/C: [nome e cargo do responsável]',after=14)
para(doc,'Assunto: Proposta de restauração e digitalização de documentos',bold=True,after=12)
para(doc,'Prezados(as),')
para(doc,'A DigitalGov, empresa com mais de 10 anos de atuação em tecnologia, apresenta à Prefeitura de Ibiá '
    'sua solução de restauração de documentos. O software recupera páginas danificadas pelo tempo e entrega, '
    'para cada página, o documento restaurado e o texto limpo transcrito, pronto para pesquisa.')

h(doc,'Como funciona')
for n,t,dsc in [('1','Envio da imagem','Foto ou digitalização de cada página, pelo celular ou computador, em lote.'),
                ('2','Restauração automática','O software remove manchas e umidade, corrige o amarelado do papel, recupera o texto apagado e melhora a nitidez.'),
                ('3','Entrega','A Prefeitura recebe o documento restaurado e o texto transcrito, com acesso seguro, busca por palavra e exportação em PDF, imagem ou texto.')]:
    p=para(doc,after=4); p.paragraph_format.left_indent=Cm(0.4)
    r=p.add_run(f'{n}. {t}: '); r.bold=True; r.font.color.rgb=NAVY
    p.add_run(dsc)

h(doc,'Preço por página')
para(doc,'DigitalGov: R$ 15 a R$ 30 por página, com o texto transcrito incluso. '
    'Referência de mercado para restauração tradicional: R$ 50 a R$ 500 por página.')

h(doc,'Exemplos por quantidade de páginas')
rows=[('1.000','R$ 15 mil – R$ 30 mil','R$ 50 mil – R$ 500 mil'),
      ('5.000','R$ 75 mil – R$ 150 mil','R$ 250 mil – R$ 2,5 milhões'),
      ('10.000','R$ 150 mil – R$ 300 mil','R$ 500 mil – R$ 5 milhões'),
      ('20.000','R$ 300 mil – R$ 600 mil','R$ 1 milhão – R$ 10 milhões'),
      ('50.000','R$ 750 mil – R$ 1,5 milhão','R$ 2,5 milhões – R$ 25 milhões')]
t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
hdr=['Páginas','DigitalGov (R$ 15–30/pág.)','Mercado (R$ 50–500/pág.)']
for i,c in enumerate(t.rows[0].cells):
    c.text=''; r=c.paragraphs[0].add_run(hdr[i]); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(0xF4,0xF7,0xFB)
    shade(c,'0E7CB8' if i==1 else '263050')
for k,row in enumerate(rows):
    cs=t.add_row().cells
    for i,v in enumerate(row):
        r=cs[i].paragraphs[0].add_run(v); r.font.size=Pt(10); r.bold=(i!=2)
        shade(cs[i],('E3F5FC' if i==1 else 'FFFFFF') if k%2==0 else ('D3EEFA' if i==1 else 'EAF0F8'))
for row in t.rows:
    for i,w in enumerate((Cm(3),Cm(6.8),Cm(6.8))): row.cells[i].width=w
para(doc,'Valores estimados (páginas × preço por página). A proposta final é definida após o diagnóstico do acervo.',
     size=9,color=GREY,italic=True,after=8).paragraph_format.space_before=Pt(4)

h(doc,'Por que a DigitalGov')
para(doc,'Mais de 10 anos de mercado, sede em Alvorada de Minas – MG, atendimento remoto, '
    'dados protegidos conforme a LGPD e acesso controlado por usuário.')
para(doc,'Colocamo-nos à disposição para uma apresentação e para o diagnóstico de uma amostra do acervo.',after=10)
para(doc,'Atenciosamente,',after=20)
para(doc,'______________________________',after=2)
para(doc,'Warley',bold=True,after=0)
para(doc,'Diretor · DigitalGov',color=GREY)
doc.save('../../deliverables/timbrado_ibia/DigitalGov_Proposta_Restauracao_Ibia.docx')
print('ok')
